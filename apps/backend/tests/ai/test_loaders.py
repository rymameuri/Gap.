from pathlib import Path

import pytest
from docx import Document
from pptx import Presentation
from pypdf import PdfWriter
from pypdf.generic import (
    DecodedStreamObject,
    DictionaryObject,
    NameObject,
)

from app.ai.rag.loaders import (
    DocumentTooLargeError,
    EmptyDocumentError,
    UnsupportedDocumentTypeError,
    load_document,
)
from app.ai.schemas.models import DocumentFormat


def create_text_pdf(path: Path) -> None:
    writer = PdfWriter()
    page = writer.add_blank_page(
        width=612,
        height=792,
    )

    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject(
                "/Helvetica"
            ),
        }
    )

    font_reference = writer._add_object(font)

    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {
                    NameObject("/F1"): font_reference,
                }
            )
        }
    )

    stream = DecodedStreamObject()
    stream.set_data(
        b"BT /F1 12 Tf 72 720 Td "
        b"(Newton laws) Tj ET"
    )

    page[NameObject("/Contents")] = writer._add_object(
        stream
    )

    with path.open("wb") as output:
        writer.write(output)


def test_loads_pdf_with_page_reference(
    tmp_path: Path,
) -> None:
    path = tmp_path / "physics.pdf"
    create_text_pdf(path)

    sections = load_document(path, "pdf-1")

    assert len(sections) == 1
    assert sections[0].format == DocumentFormat.PDF
    assert sections[0].page == 1
    assert "Newton laws" in sections[0].text


def test_loads_docx_paragraphs_and_tables(
    tmp_path: Path,
) -> None:
    path = tmp_path / "lesson.docx"

    document = Document()
    document.add_heading("Forces", level=1)
    document.add_paragraph(
        "Force equals mass times acceleration."
    )

    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Symbol"
    table.cell(0, 1).text = "Meaning"

    document.save(path)

    sections = load_document(path, "docx-1")

    assert len(sections) == 1
    assert sections[0].format == DocumentFormat.DOCX
    assert "Force equals mass times acceleration." in (
        sections[0].text
    )
    assert "Symbol | Meaning" in sections[0].text


def test_loads_pptx_with_slide_reference(
    tmp_path: Path,
) -> None:
    path = tmp_path / "lecture.pptx"

    presentation = Presentation()
    slide = presentation.slides.add_slide(
        presentation.slide_layouts[1]
    )

    slide.shapes.title.text = "Energy"
    slide.placeholders[1].text = (
        "Energy cannot be created or destroyed."
    )

    presentation.save(path)

    sections = load_document(path, "pptx-1")

    assert len(sections) == 1
    assert sections[0].format == DocumentFormat.PPTX
    assert sections[0].slide == 1
    assert "Energy" in sections[0].text


def test_rejects_unsupported_extension(
    tmp_path: Path,
) -> None:
    path = tmp_path / "notes.txt"
    path.write_text("Unsupported", encoding="utf-8")

    with pytest.raises(UnsupportedDocumentTypeError):
        load_document(path, "text-1")


def test_rejects_document_larger_than_limit(
    tmp_path: Path,
) -> None:
    path = tmp_path / "large.pdf"
    path.write_bytes(b"x" * 20)

    with pytest.raises(DocumentTooLargeError):
        load_document(
            path,
            "large-1",
            max_file_size_bytes=5,
        )


def test_rejects_empty_document(
    tmp_path: Path,
) -> None:
    path = tmp_path / "empty.docx"

    document = Document()
    document.save(path)

    with pytest.raises(EmptyDocumentError):
        load_document(path, "empty-1")