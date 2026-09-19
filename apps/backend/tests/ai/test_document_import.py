from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pptx import Presentation
from pypdf import PdfWriter

from app.ai.http_api import create_app
from app.ai.rag import document_import
from .test_http_api import FakeModelClient, FakeService, TOKEN
from .test_loaders import create_text_pdf


@pytest.fixture
def client():
    app = create_app(service=FakeService(), internal_token=TOKEN, model_client=FakeModelClient())
    with TestClient(app) as client:
        yield client


def upload(client, name, content, token=TOKEN):
    return client.post(
        "/v1/documents/extract",
        params={"filename": name},
        headers={"X-Manara-AI-Token": token, "Content-Type": "application/octet-stream"},
        content=content,
    )


@pytest.mark.parametrize("suffix, expected", [
    (".pdf", "Newton laws"),
    (".docx", "Force equals mass times acceleration"),
    (".pptx", "Energy cannot be created or destroyed"),
    (".txt", "Momentum equals mass times velocity"),
])
def test_real_document_formats(client, tmp_path: Path, suffix, expected):
    path = tmp_path / ("lecture" + suffix)
    if suffix == ".pdf":
        create_text_pdf(path)
    elif suffix == ".docx":
        document = Document()
        document.add_paragraph(expected)
        document.save(path)
    elif suffix == ".pptx":
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Energy"
        slide.placeholders[1].text = expected
        presentation.save(path)
    else:
        path.write_text(expected, encoding="utf-8-sig")
    response = upload(client, "Lecture" + suffix.upper(), path.read_bytes())
    assert response.status_code == 200, response.text
    assert expected in response.json()["text"]
    assert response.json()["section_count"] == 1
    if suffix == ".pdf":
        assert "Page 1" in response.json()["text"]
    if suffix == ".pptx":
        assert "Slide 1" in response.json()["text"]


def test_requires_internal_authentication(client):
    assert upload(client, "notes.txt", b"Physics course material", token="wrong").status_code == 401


@pytest.mark.parametrize("filename, content, code", [
    pytest.param("legacy.doc", b"Word content", "unsupported_type", id="legacy-doc"),
    pytest.param("legacy.ppt", b"PowerPoint content", "unsupported_type", id="legacy-ppt"),
    pytest.param("../course.pdf", b"bad", "invalid_document", id="parent-path"),
    pytest.param("C:\\course.pdf", b"bad", "invalid_document", id="windows-path"),
    pytest.param("bad.pdf", b"not really a PDF", "invalid_document", id="invalid-pdf"),
    pytest.param("bad.docx", b"not really a ZIP", "invalid_document", id="invalid-docx"),
    pytest.param("notes.txt", b"", "empty_document", id="empty-text"),
    pytest.param("notes.txt", b"\xff\xff", "invalid_text", id="invalid-utf8"),
    pytest.param("notes.txt", b"\x00" * 20, "invalid_text", id="nul-text"),
    pytest.param("notes.txt", b"a" * 50_001, "text_too_long", id="text-too-long"),
])
def test_rejects_unusable_uploads(client, filename, content, code):
    response = upload(client, filename, content)
    assert response.status_code == 400
    assert response.json()["detail"]["code"] == code


def test_image_only_pdf_has_a_clear_empty_error(client, tmp_path):
    path = tmp_path / "scanned.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.write(path)
    response = upload(client, path.name, path.read_bytes())
    assert response.status_code == 400
    assert response.json()["detail"]["code"] == "empty_document"


def test_encrypted_pdf(client, tmp_path):
    path = tmp_path / "locked.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.encrypt("password")
    writer.write(path)
    response = upload(client, path.name, path.read_bytes())
    assert response.json()["detail"]["code"] == "encrypted_document"


def test_enforces_streaming_size_limit(client, monkeypatch):
    monkeypatch.setattr("app.ai.http_api.MAX_UPLOAD_BYTES", 8)
    response = upload(client, "notes.txt", b"0123456789")
    assert response.status_code == 413


def test_limits_expanded_office_documents(client, tmp_path, monkeypatch):
    path = tmp_path / "large.docx"
    doc = Document()
    doc.add_paragraph("Physics lesson text")
    doc.save(path)
    monkeypatch.setattr(document_import, "MAX_EXPANDED_BYTES", 10)
    response = upload(client, path.name, path.read_bytes())
    assert response.status_code == 413


def test_counts_unicode_like_the_browser(client):
    response = upload(client, "notes.txt", ("\U0001f4d6" * 25_001).encode())
    assert response.status_code == 400
    assert response.json()["detail"]["code"] == "text_too_long"
