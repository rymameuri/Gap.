import asyncio
import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document as WordDocument

from app.ai.rag.loaders import UnsupportedDocumentTypeError
from app.ai.schemas.models import (
    AgentName,
    AgentRequest,
    CourseDocumentInput,
    CourseTextInput,
    UserRole,
)
from app.ai.service import AIService


class FakeLLMClient:
    def __init__(self) -> None:
        self.calls: list[dict[str, object]] = []

    async def acomplete(
        self,
        messages,
        *,
        max_completion_tokens: int,
    ):
        message_list = list(messages)

        self.calls.append(
            {
                "messages": message_list,
                "max_completion_tokens": max_completion_tokens,
            }
        )

        return SimpleNamespace(
            text=json.dumps(
                {
                    "answer": (
                        "Momentum is mass multiplied by "
                        "velocity [S1]."
                    ),
                    "used_source_ids": ["S1"],
                }
            )
        )


class FakeStreamingLLMClient(FakeLLMClient):
    async def astream(
        self,
        messages,
        *,
        max_completion_tokens: int,
    ):
        message_list = list(messages)
        self.calls.append(
            {
                "messages": message_list,
                "max_completion_tokens": max_completion_tokens,
                "stream": True,
            }
        )
        for delta in (
            "## Momentum\n\n",
            "Momentum is mass multiplied by velocity [S1].",
        ):
            await asyncio.sleep(0)
            yield delta


def write_docx(path: Path, text: str) -> None:
    document = WordDocument()
    document.add_heading("Course material", level=1)
    document.add_paragraph(text)
    document.save(path)


def make_request(
    course_id: str,
    message: str = "Explain momentum.",
) -> AgentRequest:
    return AgentRequest(
        user_id="student-1",
        course_id=course_id,
        role=UserRole.STUDENT,
        message=message,
    )


def test_indexes_real_word_document(
    tmp_path: Path,
) -> None:
    document_path = tmp_path / "momentum.docx"

    write_docx(
        document_path,
        (
            "Momentum is the product of an object's mass "
            "and velocity. The equation is p = mv."
        ),
    )

    service = AIService(
        llm_client=FakeLLMClient()
    )

    result = service.replace_course_documents(
        course_id="physics-101",
        documents=[
            CourseDocumentInput(
                document_id="lecture-2",
                file_path=str(document_path),
            )
        ],
    )

    assert result.course_id == "physics-101"
    assert result.document_ids == ["lecture-2"]
    assert result.document_count == 1
    assert result.section_count >= 1
    assert result.chunk_count >= 1
    assert (
        service.course_chunk_count("physics-101")
        == result.chunk_count
    )


def test_indexes_trusted_plain_text() -> None:
    service = AIService(llm_client=FakeLLMClient())

    result = service.replace_course_text(
        course_id="physics-101",
        document_id="syllabus",
        title="Physics syllabus",
        text=(
            "Momentum is the product of mass and velocity. "
            "The equation for momentum is p = mv."
        ),
    )

    assert result.document_ids == ["syllabus"]
    assert result.section_count == 1
    assert result.chunk_count >= 1
    assert service.course_chunk_count("physics-101") == result.chunk_count


def test_indexes_multiple_trusted_text_documents() -> None:
    service = AIService(llm_client=FakeLLMClient())

    result = service.replace_course_texts(
        course_id="physics-101",
        documents=[
            CourseTextInput(
                document_id="lecture-1",
                title="Lecture 1",
                text="Vectors and force are covered in lecture one.",
            ),
            CourseTextInput(
                document_id="lecture-2",
                title="Lecture 2",
                text="Momentum and collisions are covered in lecture two.",
            ),
        ],
    )

    assert result.document_ids == ["lecture-1", "lecture-2"]
    assert result.document_count == 2
    assert result.section_count == 2
    assert result.chunk_count >= 2


def test_runs_grounded_study_workflow(
    tmp_path: Path,
) -> None:
    document_path = tmp_path / "momentum.docx"

    write_docx(
        document_path,
        (
            "Momentum is the product of mass and velocity. "
            "The equation for momentum is p = mv."
        ),
    )

    llm_client = FakeLLMClient()
    service = AIService(llm_client=llm_client)

    service.replace_course_documents(
        course_id="physics-101",
        documents=[
            CourseDocumentInput(
                document_id="lecture-2",
                file_path=str(document_path),
            )
        ],
    )

    response = asyncio.run(
        service.run(make_request("physics-101"))
    )

    assert response.route.target_agent == AgentName.STUDY
    assert response.result.agent == AgentName.STUDY
    assert response.result.requires_human_review is False
    assert response.result.metadata["grounded"] is True
    assert response.result.content.endswith("[S1].")
    assert len(response.result.sources) == 1
    assert len(llm_client.calls) == 1


def test_retrieval_is_isolated_by_course(
    tmp_path: Path,
) -> None:
    course_a_path = tmp_path / "course-a.docx"
    course_b_path = tmp_path / "course-b.docx"

    write_docx(
        course_a_path,
        (
            "Momentum belongs to the PHYSICS_ALPHA course. "
            "Momentum equals mass multiplied by velocity."
        ),
    )

    write_docx(
        course_b_path,
        (
            "Database normalization belongs to the "
            "DATABASE_BETA course."
        ),
    )

    llm_client = FakeLLMClient()
    service = AIService(llm_client=llm_client)

    service.replace_course_documents(
        course_id="physics-101",
        documents=[
            CourseDocumentInput(
                document_id="physics-document",
                file_path=str(course_a_path),
            )
        ],
    )

    service.replace_course_documents(
        course_id="database-101",
        documents=[
            CourseDocumentInput(
                document_id="database-document",
                file_path=str(course_b_path),
            )
        ],
    )

    response = asyncio.run(
        service.run(make_request("physics-101"))
    )

    assert response.result.metadata["grounded"] is True

    prompt = llm_client.calls[-1]["messages"][1]["content"]

    assert "PHYSICS_ALPHA" in prompt
    assert "DATABASE_BETA" not in prompt


def test_failed_replacement_preserves_existing_index(
    tmp_path: Path,
) -> None:
    original_path = tmp_path / "original.docx"
    replacement_path = tmp_path / "replacement.docx"
    unsupported_path = tmp_path / "unsupported.txt"

    write_docx(
        original_path,
        (
            "ORIGINAL_MOMENTUM material explains that "
            "momentum equals mass multiplied by velocity."
        ),
    )

    write_docx(
        replacement_path,
        "REPLACEMENT_CONTENT for another topic.",
    )

    unsupported_path.write_text(
        "Unsupported document.",
        encoding="utf-8",
    )

    llm_client = FakeLLMClient()
    service = AIService(llm_client=llm_client)

    service.replace_course_documents(
        course_id="physics-101",
        documents=[
            CourseDocumentInput(
                document_id="original",
                file_path=str(original_path),
            )
        ],
    )

    original_chunk_count = service.course_chunk_count(
        "physics-101"
    )

    with pytest.raises(UnsupportedDocumentTypeError):
        service.replace_course_documents(
            course_id="physics-101",
            documents=[
                CourseDocumentInput(
                    document_id="replacement",
                    file_path=str(replacement_path),
                ),
                CourseDocumentInput(
                    document_id="unsupported",
                    file_path=str(unsupported_path),
                ),
            ],
        )

    assert (
        service.course_chunk_count("physics-101")
        == original_chunk_count
    )

    asyncio.run(
        service.run(make_request("physics-101"))
    )

    prompt = llm_client.calls[-1]["messages"][1]["content"]

    assert "ORIGINAL_MOMENTUM" in prompt
    assert "REPLACEMENT_CONTENT" not in prompt


def test_rejects_duplicate_document_ids(
    tmp_path: Path,
) -> None:
    document_path = tmp_path / "lecture.docx"

    write_docx(
        document_path,
        "Course material.",
    )

    service = AIService(
        llm_client=FakeLLMClient()
    )

    with pytest.raises(
        ValueError,
        match="Document IDs must be unique",
    ):
        service.replace_course_documents(
            course_id="physics-101",
            documents=[
                CourseDocumentInput(
                    document_id="lecture-1",
                    file_path=str(document_path),
                ),
                CourseDocumentInput(
                    document_id="lecture-1",
                    file_path=str(document_path),
                ),
            ],
        )


def test_full_course_request_uses_all_indexed_course_chunks() -> None:
    llm_client = FakeLLMClient()
    service = AIService(
        llm_client=llm_client,
        chunk_size_words=6,
        overlap_words=1,
    )
    index = service.replace_course_text(
        course_id="physics-101",
        document_id="syllabus",
        title="Physics syllabus",
        text=(
            "Unit one introduces force acceleration inertia. "
            "Unit two covers momentum impulse collisions. "
            "Unit three covers work energy power conservation."
        ),
    )

    response = asyncio.run(
        service.run(
            make_request(
                "physics-101",
                "Explain the full course from beginning to end.",
            )
        )
    )

    assert response.route.target_agent == AgentName.STUDY
    assert response.result.metadata["scope"] == "course_overview"
    assert response.result.metadata["retrieved_count"] == index.chunk_count
    assert llm_client.calls[-1]["max_completion_tokens"] == 3_000
    prompt = llm_client.calls[-1]["messages"][1]["content"]
    for term in (
        "force",
        "acceleration",
        "momentum",
        "impulse",
        "collisions",
        "work",
        "energy",
        "conservation",
    ):
        assert term in prompt

def test_stream_study_preserves_review_gate_and_progressive_content() -> None:
    llm_client = FakeStreamingLLMClient()
    service = AIService(llm_client=llm_client)
    service.replace_course_text(
        course_id="physics-101",
        document_id="syllabus",
        title="Physics syllabus",
        text="Momentum is the product of mass and velocity.",
    )
    streamed: list[str] = []

    async def run():
        return await service.stream_study(
            make_request("physics-101"),
            lambda text: _append(streamed, text),
        )

    result = asyncio.run(run())

    assert result.requires_human_review is False
    assert result.metadata["review_approved"] is True
    assert "".join(streamed) == result.content
    assert result.content.endswith("[S1].")
    assert llm_client.calls[-1]["stream"] is True


async def _append(values: list[str], text: str) -> None:
    values.append(text)

