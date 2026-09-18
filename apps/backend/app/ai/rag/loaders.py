from __future__ import annotations

from pathlib import Path

from docx import Document
from pypdf import PdfReader
from pptx import Presentation

from app.ai.schemas.models import (
    DocumentFormat,
    DocumentSection,
)


DEFAULT_MAX_FILE_SIZE_BYTES = 25 * 1024 * 1024

SUPPORTED_EXTENSIONS = {
    ".pdf": DocumentFormat.PDF,
    ".docx": DocumentFormat.DOCX,
    ".pptx": DocumentFormat.PPTX,
}


class DocumentLoadError(RuntimeError):
    """Base error for document loading failures."""


class UnsupportedDocumentTypeError(DocumentLoadError):
    """Raised when the document extension is unsupported."""


class DocumentTooLargeError(DocumentLoadError):
    """Raised when a document exceeds the configured limit."""


class EmptyDocumentError(DocumentLoadError):
    """Raised when no readable text is found."""


def _normalise_text(text: str) -> str:
    lines = []

    for raw_line in text.replace("\xa0", " ").splitlines():
        line = " ".join(raw_line.split())

        if line:
            lines.append(line)

    return "\n".join(lines)


def _load_pdf(
    path: Path,
    document_id: str,
) -> list[DocumentSection]:
    try:
        reader = PdfReader(str(path))

        if reader.is_encrypted:
            raise DocumentLoadError(
                "Password-protected PDFs are not supported."
            )

        sections = []

        for page_number, page in enumerate(
            reader.pages,
            start=1,
        ):
            text = _normalise_text(page.extract_text() or "")

            if not text:
                continue

            sections.append(
                DocumentSection(
                    document_id=document_id,
                    title=path.stem,
                    format=DocumentFormat.PDF,
                    section_index=page_number,
                    page=page_number,
                    text=text,
                )
            )

        return sections
    except DocumentLoadError:
        raise
    except Exception as exc:
        raise DocumentLoadError(
            f"Could not read PDF: {path.name}"
        ) from exc


def _load_docx(
    path: Path,
    document_id: str,
) -> list[DocumentSection]:
    try:
        document = Document(str(path))
        text_parts = []

        for paragraph in document.paragraphs:
            text = _normalise_text(paragraph.text)

            if text:
                text_parts.append(text)

        for table in document.tables:
            for row in table.rows:
                cells = [
                    _normalise_text(cell.text)
                    for cell in row.cells
                ]

                row_text = " | ".join(
                    cell for cell in cells if cell
                )

                if row_text:
                    text_parts.append(row_text)

        combined_text = "\n".join(text_parts)

        if not combined_text:
            return []

        return [
            DocumentSection(
                document_id=document_id,
                title=path.stem,
                format=DocumentFormat.DOCX,
                section_index=1,
                text=combined_text,
            )
        ]
    except Exception as exc:
        raise DocumentLoadError(
            f"Could not read Word document: {path.name}"
        ) from exc


def _load_pptx(
    path: Path,
    document_id: str,
) -> list[DocumentSection]:
    try:
        presentation = Presentation(str(path))
        sections = []

        for slide_number, slide in enumerate(
            presentation.slides,
            start=1,
        ):
            text_parts = []

            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    text = _normalise_text(shape.text)

                    if text:
                        text_parts.append(text)

                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        cells = [
                            _normalise_text(cell.text)
                            for cell in row.cells
                        ]

                        row_text = " | ".join(
                            cell for cell in cells if cell
                        )

                        if row_text:
                            text_parts.append(row_text)

            combined_text = "\n".join(text_parts)

            if not combined_text:
                continue

            sections.append(
                DocumentSection(
                    document_id=document_id,
                    title=path.stem,
                    format=DocumentFormat.PPTX,
                    section_index=slide_number,
                    slide=slide_number,
                    text=combined_text,
                )
            )

        return sections
    except Exception as exc:
        raise DocumentLoadError(
            f"Could not read PowerPoint: {path.name}"
        ) from exc


def load_document(
    file_path: str | Path,
    document_id: str,
    max_file_size_bytes: int = DEFAULT_MAX_FILE_SIZE_BYTES,
) -> list[DocumentSection]:
    path = Path(file_path)

    if not document_id.strip():
        raise ValueError("document_id cannot be empty.")

    if max_file_size_bytes < 1:
        raise ValueError(
            "max_file_size_bytes must be greater than zero."
        )

    if not path.exists() or not path.is_file():
        raise DocumentLoadError(
            f"Document does not exist: {path}"
        )

    if path.is_symlink():
        raise DocumentLoadError(
            "Symbolic links are not accepted as documents."
        )

    document_format = SUPPORTED_EXTENSIONS.get(
        path.suffix.casefold()
    )

    if document_format is None:
        raise UnsupportedDocumentTypeError(
            "Supported document types are PDF, DOCX, and PPTX."
        )

    file_size = path.stat().st_size

    if file_size > max_file_size_bytes:
        raise DocumentTooLargeError(
            f"Document size {file_size} exceeds the configured limit."
        )

    loader = {
        DocumentFormat.PDF: _load_pdf,
        DocumentFormat.DOCX: _load_docx,
        DocumentFormat.PPTX: _load_pptx,
    }[document_format]

    sections = loader(path, document_id)

    if not sections:
        raise EmptyDocumentError(
            "No readable text was found in the document."
        )

    return sections